"""Executive report and Excel workpaper generation for the benchmark."""

from __future__ import annotations

from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src import config
from src.benchmarking import run_benchmark
from src.comparables import compute_cascade_stats, get_accepted, get_rejected
from src.pli_calculator import yearly_pli_table
from src.sensitivity import run_all_scenarios, scenario_summary_frame

REPORT_FILENAME = "pfizer_tp_benchmark_workpaper.xlsx"
DECISIONS_FILENAME = "comparables_decisions.xlsx"
REPORT_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_report_context(
    tested_party_df: pd.DataFrame,
    raw_comparables_df: pd.DataFrame,
    decisions_df: pd.DataFrame,
    tested_party_name: str | None = None,
) -> dict[str, Any]:
    """Build the shared data context for the report page and Excel export.

    Args:
        tested_party_df: Tested-party Orbis data.
        raw_comparables_df: Raw Orbis comparable-candidate data.
        decisions_df: Comparable-candidate decision state.
        tested_party_name: Optional display name for tested-party labels.

    Returns:
        Dictionary with base benchmark results, cascade data, and sensitivity
        outputs used by both Streamlit and Excel.
    """

    display_name = tested_party_name or config.tested_party_display_name()
    accepted = get_accepted(raw_comparables_df, decisions_df)
    rejected = get_rejected(raw_comparables_df, decisions_df)
    base_result = run_benchmark(
        tested_party_df=tested_party_df,
        comparables_df=accepted,
        pli_type=config.PLI_OPERATING_MARGIN,
        years=config.DEFAULT_BENCHMARK_PERIOD,
    )
    scenarios = run_all_scenarios(
        tested_party_df=tested_party_df,
        comparables_df=accepted,
        include_fy22_adjustment=True,
        tested_party_name=display_name,
    )
    return {
        "tested_party_name": display_name,
        "tested_party": tested_party_df,
        "raw_comparables": raw_comparables_df,
        "decisions": decisions_df,
        "accepted": accepted,
        "rejected": rejected,
        "cascade_stats": compute_cascade_stats(
            decisions_df,
            raw_count=len(raw_comparables_df),
        ),
        "base_result": base_result,
        "scenarios": scenarios,
        "sensitivity_summary": scenario_summary_frame(scenarios),
    }


def executive_conclusion(context: dict[str, Any]) -> str:
    """Return a concise executive conclusion for the report page.

    Args:
        context: Report context from `build_report_context`.

    Returns:
        Markdown-ready conclusion paragraph.
    """

    result = context["base_result"]
    range_dict = result["range"]
    stats = context["cascade_stats"]
    position = result["position"]["position"]
    conclusion = _position_phrase(str(position))
    period = result["period_label"]
    tested_pli = _format_percent(result["tested_pli"])
    q1 = _format_percent(range_dict["q1"])
    median = _format_percent(range_dict["median"])
    q3 = _format_percent(range_dict["q3"])
    tested_party_name = context["tested_party_name"]

    sensitivity = context["sensitivity_summary"]
    changed = sensitivity.loc[sensitivity["Position"] != position, "Scenario"].tolist()
    if changed:
        sensitivity_sentence = (
            "Sensitivity checks are not unanimous: "
            + "; ".join(changed[:3])
            + " change the tested-party position."
        )
    else:
        sensitivity_sentence = (
            "Sensitivity checks do not change the base-case position."
        )

    return (
        f"{tested_party_name} is characterized as an LRD-SM and tested under "
        f"TNMM using Operating Margin as the primary PLI over {period}. "
        f"The rejection cascade narrows the Orbis candidate pool from "
        f"{stats['raw']} companies to {stats['accepted']} accepted comparables. "
        f"{tested_party_name}'s weighted Operating Margin is {tested_pli}, "
        f"compared with an "
        f"interquartile range of {q1} to {q3} and median of {median}. "
        f"On this base case, the tested party is **{conclusion}**. "
        f"{sensitivity_sentence} The result should be read with the documented "
        f"limitations: the accepted pool is small, {tested_party_name} is "
        "materially larger than the median comparable, and independent "
        "multinational pharma distributors are scarce in Europe."
    )


def overview_frame(context: dict[str, Any]) -> pd.DataFrame:
    """Return a compact overview table for display and Excel.

    Args:
        context: Report context from `build_report_context`.

    Returns:
        Two-column overview DataFrame.
    """

    tested_party = context["tested_party"]
    tested_party_name = context["tested_party_name"]
    result = context["base_result"]
    stats = context["cascade_stats"]
    row = tested_party.iloc[0]
    nace = _format_nace(row.get(config.NACE_COLUMN))
    return pd.DataFrame(
        [
            ("Tested party", row.get(config.COMPANY_NAME_COLUMN, tested_party_name)),
            ("Country", row.get(config.COUNTRY_COLUMN, "Germany")),
            ("NACE", nace),
            ("Characterization", "Limited-Risk Distributor with Sales and Marketing"),
            ("Primary PLI", result["pli_label"]),
            ("Benchmark period", result["period_label"]),
            ("Raw candidates", stats["raw"]),
            ("Accepted comparables", stats["accepted"]),
            ("Rejected candidates", stats["rejected"]),
            ("Base-case conclusion", _position_phrase(result["position"]["position"])),
        ],
        columns=["Item", "Value"],
    )


def range_frame(context: dict[str, Any]) -> pd.DataFrame:
    """Return the arm's-length range summary table.

    Args:
        context: Report context from `build_report_context`.

    Returns:
        Arm's-length range metrics with raw decimal values.
    """

    range_dict = context["base_result"]["range"]
    return pd.DataFrame(
        [
            ("N comparables", range_dict["n"]),
            ("Minimum", range_dict["min"]),
            ("Q1 (25th percentile)", range_dict["q1"]),
            ("Median", range_dict["median"]),
            ("Q3 (75th percentile)", range_dict["q3"]),
            ("Maximum", range_dict["max"]),
            ("IQR width", range_dict["iqr_width"]),
            (
                f"{context['tested_party_name']} weighted OM",
                context["base_result"]["tested_pli"],
            ),
        ],
        columns=["Metric", "Value"],
    )


def accepted_comparables_frame(context: dict[str, Any]) -> pd.DataFrame:
    """Return accepted comparables with base-case PLI detail.

    Args:
        context: Report context from `build_report_context`.

    Returns:
        Accepted comparable-company table for reporting.
    """

    detail = context["base_result"]["comparables_detail"].copy()
    detail["Revenue EURm"] = detail[config.LATEST_REVENUE_COLUMN] / 1_000
    return detail[
        [
            config.COMPANY_NAME_COLUMN,
            config.COUNTRY_COLUMN,
            "Revenue EURm",
            config.LATEST_EMPLOYEES_COLUMN,
            "weighted_pli",
            "inside_iqr",
            config.TRADE_DESCRIPTION_COLUMN,
        ]
    ].rename(
        columns={
            config.COMPANY_NAME_COLUMN: "Company",
            config.COUNTRY_COLUMN: "Country",
            config.LATEST_EMPLOYEES_COLUMN: "Employees",
            "weighted_pli": "Weighted Operating Margin",
            "inside_iqr": "Inside IQR",
            config.TRADE_DESCRIPTION_COLUMN: "Trade description",
        }
    )


def rejected_candidates_frame(context: dict[str, Any]) -> pd.DataFrame:
    """Return rejected candidates with category labels and notes.

    Args:
        context: Report context from `build_report_context`.

    Returns:
        Rejected candidate table for reporting.
    """

    rejected = context["rejected"].copy()
    rejected["Revenue EURm"] = rejected[config.LATEST_REVENUE_COLUMN] / 1_000
    return rejected[
        [
            config.COMPANY_NAME_COLUMN,
            config.COUNTRY_COLUMN,
            "reason_label",
            "notes",
            "Revenue EURm",
            config.TRADE_DESCRIPTION_COLUMN,
        ]
    ].rename(
        columns={
            config.COMPANY_NAME_COLUMN: "Company",
            config.COUNTRY_COLUMN: "Country",
            "reason_label": "Rejection category",
            "notes": "Methodology note",
            config.TRADE_DESCRIPTION_COLUMN: "Trade description",
        }
    )


def pli_detail_frame(context: dict[str, Any]) -> pd.DataFrame:
    """Return yearly and weighted PLI detail for accepted comparables.

    Args:
        context: Report context from `build_report_context`.

    Returns:
        Comparable-company yearly PLI table.
    """

    return yearly_pli_table(
        context["accepted"],
        config.PLI_OPERATING_MARGIN,
        config.DEFAULT_BENCHMARK_PERIOD,
    ).rename(columns={config.COMPANY_NAME_COLUMN: "Company"})


def sensitivity_frame(context: dict[str, Any]) -> pd.DataFrame:
    """Return scenario summary for the report.

    Args:
        context: Report context from `build_report_context`.

    Returns:
        Sensitivity scenario table.
    """

    summary = context["sensitivity_summary"].copy()
    base_position = context["base_result"]["position"]["position"]
    summary["Conclusion changes"] = summary["Position"] != base_position
    return summary


def methodology_notes_frame(
    tested_party_name: str | None = None,
) -> pd.DataFrame:
    """Return methodology notes for the Excel report.

    Returns:
        Methodology notes DataFrame.
    """

    display_name = tested_party_name or config.tested_party_display_name()
    notes = [
        (
            "Tested party",
            f"{display_name} is characterized as an LRD-SM based on the FAR memo.",
        ),
        (
            "Primary method",
            "TNMM compares net profitability of the tested party with independent "
            "companies performing broadly comparable distribution functions.",
        ),
        (
            "Primary PLI",
            "Operating Margin is used as the primary PLI; Berry Ratio and ROCE "
            "are retained as secondary sensitivity checks.",
        ),
        (
            "Multi-year weighting",
            "Operating Margin is calculated as sum(EBIT) / sum(Sales) over the "
            "selected period, not as an average of yearly ratios.",
        ),
        (
            "Range construction",
            f"The arm's-length range uses the interquartile range with "
            f"{config.QUARTILE_METHOD} quartiles.",
        ),
        (
            "Limitations",
            f"The accepted pool is small and {display_name} is materially larger "
            "than the accepted comparable pool; results should be read with this "
            "caveat.",
        ),
    ]
    return pd.DataFrame(notes, columns=["Topic", "Note"])


def build_decisions_workbook(decisions_df: pd.DataFrame) -> bytes:
    """Build an Excel workbook for comparable-decision review.

    Args:
        decisions_df: Comparable-candidate decision state.

    Returns:
        XLSX workbook bytes with decisions and rejection-category notes.
    """

    category_rows = [
        {
            "reason": reason,
            "label": metadata["label"],
            "description": metadata["description"],
        }
        for reason, metadata in config.REJECT_CATEGORIES.items()
    ]

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        decisions_df.to_excel(writer, sheet_name="Decisions", index=False)
        pd.DataFrame(category_rows).to_excel(
            writer,
            sheet_name="Category Notes",
            index=False,
        )
        _format_workbook(writer.book)

    output.seek(0)
    return output.read()


def build_excel_report(
    tested_party_df: pd.DataFrame,
    raw_comparables_df: pd.DataFrame,
    decisions_df: pd.DataFrame,
    tested_party_name: str | None = None,
) -> bytes:
    """Build the downloadable Excel workpaper.

    Args:
        tested_party_df: Tested-party Orbis data.
        raw_comparables_df: Raw Orbis comparable-candidate data.
        decisions_df: Comparable-candidate decision state.
        tested_party_name: Optional display name for tested-party labels.

    Returns:
        XLSX workbook bytes.
    """

    context = build_report_context(
        tested_party_df,
        raw_comparables_df,
        decisions_df,
        tested_party_name=tested_party_name,
    )
    sheets = {
        "Overview": overview_frame(context),
        "Accepted Comparables": accepted_comparables_frame(context),
        "Rejected Candidates": rejected_candidates_frame(context),
        "PLI Detail": pli_detail_frame(context),
        "Arm's-Length Range": range_frame(context),
        "Sensitivity Scenarios": sensitivity_frame(context),
        "Methodology Notes": methodology_notes_frame(context["tested_party_name"]),
    }

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, dataframe in sheets.items():
            dataframe.to_excel(writer, sheet_name=sheet_name, index=False)
        _format_workbook(writer.book)

    output.seek(0)
    return output.read()


def _format_workbook(workbook: Any) -> None:
    """Apply light formatting to all report worksheets."""

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for column_cells in worksheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            width = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            worksheet.column_dimensions[column_letter].width = min(
                max(width + 2, 12), 48
            )
        _format_percent_columns(worksheet)


def _format_percent_columns(worksheet: Any) -> None:
    """Apply percentage number format to PLI and range columns."""

    percent_headers = {
        "Weighted Operating Margin",
        "Weighted PLI",
        "FY 2024",
        "FY 2023",
        "FY 2022",
        "Tested PLI",
        "Q1",
        "Median",
        "Q3",
    }
    for cell in worksheet[1]:
        if cell.value in percent_headers:
            for row_cell in worksheet.iter_cols(
                min_col=cell.column,
                max_col=cell.column,
                min_row=2,
            ):
                for value_cell in row_cell:
                    value_cell.number_format = "0.00%"


def _format_percent(value: object) -> str:
    """Format a decimal PLI as a percent string."""

    if pd.isna(value):
        return "n/a"
    return f"{float(value) * 100:.2f}%"


def _format_nace(value: object) -> str:
    """Format NACE code for reporting."""

    if pd.isna(value):
        return "n/a"
    try:
        code = f"{int(float(value)):04d}"
    except (TypeError, ValueError):
        code = str(value)
    description = config.NACE_DESCRIPTIONS.get(code)
    return f"{code} - {description}" if description else code


def _position_phrase(position: object) -> str:
    """Return a human-readable conclusion phrase."""

    labels = {
        "below_q1": "below the arm's-length range",
        "within_range": "within the arm's-length range",
        "above_q3": "above the arm's-length range",
        "not_available": "not available",
    }
    return labels.get(str(position), str(position))
