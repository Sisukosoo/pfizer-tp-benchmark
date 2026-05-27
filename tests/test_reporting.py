"""Tests for executive report and Excel workpaper generation."""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from src import config
from src.default_decisions import make_decision_key
from src.reporting import (
    build_excel_report,
    build_report_context,
    executive_conclusion,
)


def test_build_report_context_accepts_decision_pool(
    sample_tested_party_df,
    sample_comparables_df,
) -> None:
    """Report context should attach accepted comparables and benchmark output."""

    decisions = _accept_all_decisions(sample_comparables_df)
    context = build_report_context(
        sample_tested_party_df,
        sample_comparables_df,
        decisions,
    )

    assert len(context["accepted"]) == 3
    assert context["base_result"]["range"]["n"] == 3
    assert "arm's-length range" in executive_conclusion(context)


def test_excel_report_contains_expected_sheets(
    sample_tested_party_df,
    sample_comparables_df,
) -> None:
    """Generated workbook should include the expected workpaper tabs."""

    decisions = _accept_all_decisions(sample_comparables_df)
    report_bytes = build_excel_report(
        sample_tested_party_df,
        sample_comparables_df,
        decisions,
    )
    workbook = load_workbook(BytesIO(report_bytes), read_only=True)

    assert report_bytes.startswith(b"PK")
    assert workbook.sheetnames == [
        "Overview",
        "Accepted Comparables",
        "Rejected Candidates",
        "PLI Detail",
        "Arm's-Length Range",
        "Sensitivity Scenarios",
        "Methodology Notes",
    ]


def _accept_all_decisions(sample_comparables_df):
    """Return accept decisions for the synthetic comparable pool."""

    rows = []
    for _, row in sample_comparables_df.iterrows():
        rows.append(
            {
                "bvd_id": make_decision_key(
                    row[config.COMPANY_NAME_COLUMN],
                    row[config.COUNTRY_COLUMN],
                ),
                "company_name": row[config.COMPANY_NAME_COLUMN],
                "country": row[config.COUNTRY_COLUMN],
                "decision": config.DECISION_ACCEPT,
                "reason": "",
                "notes": "",
                "modified_at": "2026-05-27T00:00:00+00:00",
                "modified_by": "test",
            }
        )
    return sample_comparables_df.__class__(rows)
